"""Placeholder filling in Word templates: the cases Word and users actually produce."""
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import labels

ROW = {"family": "Rosaceae", "species": "Rosa canina L.", "date": datetime(2025, 6, 3), "num": 101,
       "N": 56.0153, "E": 37.8471, "inat": 281234567, "Altitude": None, "habitats": "near {species} sign"}


def render(*paragraphs, table=None):
    """Template made of paragraphs given as lists of (text, bold) runs; returns (lines, pictures)."""
    df = pd.DataFrame([ROW])
    columns = labels.resolve_columns(df.columns)
    values = [labels.row_values(r, columns, "AR", "Vascular plant", "google", True) for _, r in df.iterrows()]
    doc = Document()
    for pieces in paragraphs:
        p = doc.add_paragraph()
        for text, bold in pieces:
            p.add_run(text).bold = bold
    if table:
        doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(table)
    cell = labels.render(doc, values, 1).tables[0].rows[0].cells[0]
    lines = [Paragraph(p, cell).text for p in cell._tc.iter(qn("w:p"))]
    return [line for line in lines if line], len(cell._tc.xpath(".//pic:pic"))


def run():
    plain = lambda text: [(text, False)]
    assert render([("{spe", False), ("cies}", True)]) == (["Rosa canina L."], 0)  # Word split the run
    assert render([("Sp: {", False), ("speci", True), ("es} ok", False)])[0] == ["Sp: Rosa canina L. ok"]
    assert render(plain("{ species }"))[0] == ["Rosa canina L."]
    assert render(plain("{Species}"))[0] == ["Rosa canina L."]  # capital letters don't matter
    assert render(plain("{date} / № {num}"))[0] == ["2025-06-03 / № 101"]
    assert render([("QR: {q", False), ("r} end", False)]) == (["QR:  end"], 1)
    assert render(plain("{qr}{inat_qr}"))[1] == 2
    assert "№ 101" in render(table="№ {num}")[0]
    assert render(plain("{foo}"))[0] == ["{foo}"]  # unknown stays visible
    assert render(plain("{habitats}"))[0] == ["near {species} sign"]  # data is not expanded again
    assert render(plain("{species"))[0] == ["{species"]

    # Empty values
    assert render(plain("{species}"), plain("{Altitude}"))[0] == ["Rosa canina L."]
    assert render(plain("{species}"), plain("Alt.: {Altitude}"))[0] == ["Rosa canina L."]
    assert render(plain("Det.: {Altitude} ____"))[0] == ["Det.:  ____"]  # kept for handwriting
    assert render(plain("Alt.: {Altitude} / № {num}"))[0] == ["Alt.:  / № 101"]
    assert render(plain("{family} ({Altitude})"))[0] == ["Rosaceae"]
    assert render([("{family} (", False), ("{num}", True), (")", False)])[0] == ["Rosaceae (101)"]

    template = Document()
    template.add_paragraph("{Species} {foo} {ALTITUDE} {Bar}")
    assert labels.unknown_placeholders(template, pd.DataFrame([ROW]).columns) == ["Bar", "foo"]
    # Preview: the filled label as HTML, with the QR codes as images and the label width
    df = pd.DataFrame([ROW])
    values = labels.row_values(df.iloc[0], labels.resolve_columns(df.columns), "AR", "Vascular plant", "osm", True)
    html, images, width = labels.preview_html(labels.default_template(), values, 2)
    assert "Rosa canina L." in html and "56.0153° N" in html and "{qr}" not in html, html
    assert len(images) == 2 and width == 340  # 18 cm of printable width / 2 at 96 dpi
    assert labels.preview_html(labels.default_template(), values, 3)[2] == 227
    print("labels test passed")


if __name__ == "__main__":
    run()
