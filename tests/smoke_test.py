"""End-to-end check: Excel -> labels .docx, without showing any window."""
import os
import sys
import tempfile
from datetime import datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import pandas as pd
from docx import Document
from PyQt6.QtCore import QSettings
from PyQt6.QtWidgets import QApplication, QFileDialog, QMessageBox

import labels
import main


def run():
    tmp = tempfile.mkdtemp()
    xlsx = os.path.join(tmp, "in.xlsx")
    pd.DataFrame({
        "family": ["Rosaceae", "Poaceae", "Fabaceae"],
        "familyrus": ["Розоцветные", None, "Бобовые"],
        "species": ["Rosa canina", "Poa annua", "Trifolium repens"],
        "speciesrus": ["Шиповник", "Мятлик", None],
        "region": ["Moscow reg.", None, "Tver reg."],
        "point": ["Pushkino", "Dubna", None],
        "habitats": ["forest edge", None, "meadow"],
        "date": [datetime(2025, 6, 1, 10, 0), datetime(2025, 6, 2), "2025-06-03"],
        "leg.": ["A. S."] * 3,
        "det.": ["A. S."] * 3,
        "num": [12, None, 14],
        "N": ["55,75", 0, None],
        "E": [37.6, 0.0, 36.1],
        "inat": [123456, None, None],
        "Altitude": ["150 m", None, None],
    }).to_excel(xlsx, index=False)

    # Keep the test away from the real user settings
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, tmp)
    app = QApplication([])
    errors = []
    QMessageBox.information = lambda *a: None
    QMessageBox.critical = lambda *a: errors.append(a[2])
    QMessageBox.warning = lambda *a: errors.append(a[2])

    window = main.HerbariumApp()
    window.file_path = xlsx
    window.herbarium_input.setText("AR Herbarium")
    window.set_herbarium_type("Vascular plant")
    window.qr_checkbox.setChecked(True)
    window.process_file()
    assert not errors, errors

    table = Document(os.path.join(tmp, "in_output.docx")).tables[0]
    # Labels fill the grid row by row
    rosa, poa = table.rows[0].cells
    trifolium = table.rows[1].cells[0]
    assert rosa.text.startswith("AR Herbarium\nRosaceae\nРозоцветные"), rosa.text
    assert "№ 12\t" in rosa.text and "2025-06-01\t" in rosa.text and "55.75° N, 37.6° E" in rosa.text, rosa.text
    assert ", ," not in poa.text and "Dubna" in poa.text, poa.text
    assert "Мятлик" in poa.text and "\n\n" not in poa.text.strip(), poa.text  # empty familyrus line removed
    images = lambda cell: len(cell._element.xpath(".//pic:pic"))
    assert images(rosa) == 1 and images(poa) == 1  # "55,75" and 0/0 both give a QR
    assert images(trifolium) == 0  # no latitude -> no QR

    # iNaturalist QR, three labels per row and a user template with a custom column
    template = labels.default_template()
    template.add_paragraph("Alt.: {Altitude}")
    template.save(os.path.join(tmp, "template.docx"))
    window.inat_checkbox.setChecked(True)
    main.set_segment(window.template_group, "custom")
    window.template_path = os.path.join(tmp, "template.docx")
    main.set_segment(window.per_row_group, 3)
    window.process_file()
    assert not errors, errors
    row = Document(os.path.join(tmp, "in_output.docx")).tables[0].rows[0].cells
    assert len(row) == 3 and images(row[0]) == 2 and "Alt.: 150 m" in row[0].text, row[0].text

    # Darwin Core export
    QFileDialog.getSaveFileName = lambda *a: (os.path.join(tmp, "dwc.csv"), "")
    window.export_darwin_core()
    dwc = pd.read_csv(os.path.join(tmp, "dwc.csv"))
    assert list(dwc["catalogNumber"].fillna(0).astype(int)) == [12, 0, 14] and dwc["decimalLatitude"][0] == 55.75

    for mode in main.THEME_MODES:  # every theme stylesheet must format cleanly
        window.cycle_theme()
    assert main.parse_version("1.10") > main.parse_version("1.9")
    app.quit()
    print("smoke test passed")


if __name__ == "__main__":
    run()
