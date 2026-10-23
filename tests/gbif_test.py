"""GBIF name matching and Darwin Core export, on recorded GBIF answers (no network)."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

import pandas as pd
from openpyxl import Workbook, load_workbook

import gbif
import labels

ROSA = {"usage": {"name": "Rosa canina L.", "canonicalName": "Rosa canina", "rank": "SPECIES", "status": "ACCEPTED"},
        "classification": [{"name": "Rosaceae", "rank": "FAMILY"}], "diagnostics": {"matchType": "EXACT"}}
TYPO = {**ROSA, "diagnostics": {"matchType": "VARIANT"}}
SYNONYM = {"usage": {"name": "Pyrethrum parthenium (L.) Sm.", "canonicalName": "Pyrethrum parthenium",
                     "rank": "SPECIES", "status": "SYNONYM"},
           "acceptedUsage": {"name": "Tanacetum parthenium (L.) Sch.Bip.", "rank": "SPECIES"},
           "classification": [{"name": "Asteraceae", "rank": "FAMILY"}], "diagnostics": {"matchType": "EXACT"}}
GENUS_ONLY = {"usage": {"name": "Rosa L.", "rank": "GENUS"}, "diagnostics": {"matchType": "HIGHERRANK"}}
NOTHING = {"usage": {"name": "Plantae", "rank": "KINGDOM"}, "diagnostics": {"matchType": "HIGHERRANK"}}


def run():
    assert gbif.interpret_match("Rosa canina L.", ROSA) == (None, "Rosaceae", None)
    suggestion, _, _ = gbif.interpret_match("Rosa canina", ROSA)
    assert suggestion.new == "Rosa canina L." and not suggestion.apply  # author only: offered, not pre-ticked
    suggestion, _, _ = gbif.interpret_match("Rosa caninna", TYPO)
    assert suggestion.new == "Rosa canina L." and suggestion.apply
    suggestion, family, _ = gbif.interpret_match("Pyrethrum parthenium", SYNONYM)
    assert suggestion.new == "Tanacetum parthenium (L.) Sch.Bip." and family == "Asteraceae"
    assert "only Rosa L." in gbif.interpret_match("Rosa xyz", GENUS_ONLY)[2]
    assert gbif.interpret_match("Foo bar", NOTHING)[0] is None

    assert gbif.family_suggestion("Asteracaea", "Asteraceae").apply
    assert not gbif.family_suggestion("Compositae", "Asteraceae").apply
    assert gbif.family_suggestion("Rosaceae", "Rosaceae") is None

    # A family name shared by species GBIF puts in different families is reported, never renamed
    checker = gbif.NameChecker.__new__(gbif.NameChecker)
    checker.problems = []
    checker.family_votes = {"Rosaceae": {"Rosaceae": ["Rosa canina"], "Poaceae": ["Poa annua"]},
                            "Asteracaea": {"Asteraceae": ["Achillea millefolium"]}}
    suggestions = checker.family_suggestions()
    assert [(s.old, s.new) for s in suggestions] == [("Asteracaea", "Asteraceae")]
    assert checker.problems == [("Poa annua", "Family Rosaceae in the table, GBIF says Poaceae")]

    # Corrections go into a copy of the workbook and keep the other cells
    folder = tempfile.mkdtemp()
    source, target = os.path.join(folder, "in.xlsx"), os.path.join(folder, "out.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["family", "species", "num"])
    sheet.append(["Rosaceae", "Rosa caninna ", 1])
    sheet.append(["Asteracaea", "Achillea millefolium", 2])
    workbook.save(source)
    columns = labels.resolve_columns(["family", "species", "num"])
    changed = gbif.save_corrected_copy(source, target, columns, [
        gbif.Suggestion("species", "Rosa caninna", "Rosa canina L.", "", True),
        gbif.Suggestion("family", "Asteracaea", "Asteraceae", "", True)])
    rows = list(load_workbook(target).active.values)
    assert changed == 2 and rows[1] == ("Rosaceae", "Rosa canina L.", 1) and rows[2][0] == "Asteraceae", rows
    assert list(load_workbook(source).active.values)[1][1] == "Rosa caninna "  # original untouched

    df = pd.DataFrame({"species": ["Rosa canina L."], "family": ["Rosaceae"], "num": [12.0], "N": ["55,75"],
                       "E": [37.6], "date": [pd.Timestamp("2025-06-01 10:00")], "inat": [123]})
    dwc = gbif.to_darwin_core(df, labels.resolve_columns(df.columns), "AR", "Vascular plant").iloc[0]
    assert dwc["occurrenceID"] == "AR:12" and dwc["decimalLatitude"] == 55.75 and dwc["eventDate"] == "2025-06-01"
    assert dwc["kingdom"] == "Plantae" and dwc["associatedReferences"].endswith("/observations/123")
    print("gbif test passed")


if __name__ == "__main__":
    run()
