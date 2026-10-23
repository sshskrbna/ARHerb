"""GBIF: checking scientific names against the GBIF Backbone and Darwin Core export."""
import difflib
import json
from dataclasses import dataclass
from urllib.parse import urlencode

import pandas as pd
from openpyxl import load_workbook
from PyQt6.QtCore import QObject, QUrl, pyqtSignal
from PyQt6.QtNetwork import QNetworkAccessManager, QNetworkReply, QNetworkRequest

from labels import filter_nan, format_num, inat_url, parse_coordinate, remove_time_from_date

MATCH_URL = "https://api.gbif.org/v2/species/match"


@dataclass
class Suggestion:
    field: str        # "species" or "family"
    old: str
    new: str
    reason: str
    apply: bool       # pre-ticked in the dialog


def kingdom_for(herbarium_type):
    return "Fungi" if herbarium_type == "Fungi" else "Plantae"


def interpret_match(name, response):
    """Turn a GBIF v2 match response into (species suggestion or None, GBIF family or None, problem text or None)."""
    usage = response.get("usage") or {}
    accepted = response.get("acceptedUsage") or {}
    match_type = (response.get("diagnostics") or {}).get("matchType", "NONE")
    family = next((c["name"] for c in response.get("classification", []) if c.get("rank") == "FAMILY"), None)

    if match_type == "NONE" or not usage or usage.get("rank") not in ("SPECIES", "SUBSPECIES", "VARIETY", "FORM"):
        found = f" (only {usage['name']} was found)" if usage.get("rank") in ("GENUS", "FAMILY") else ""
        return None, None, f"Not found in GBIF{found}"

    if accepted:
        return Suggestion("species", name, accepted["name"], f"Synonym of {accepted['name']}", True), family, None
    if match_type in ("VARIANT", "FUZZY"):
        return Suggestion("species", name, usage["name"], "Probably a typo", True), family, None
    if name != usage["name"]:
        # Same species, different or missing author: worth showing, not worth forcing
        reason = "Add the author" if name == usage.get("canonicalName") else "Author differs"
        return Suggestion("species", name, usage["name"], reason, False), family, None
    return None, family, None


def family_suggestion(old, new):
    if not old or not new or old == new:
        return None
    similar = difflib.SequenceMatcher(None, old.lower(), new.lower()).ratio() > 0.8
    # Very different names are often valid alternatives (Compositae / Asteraceae), so not pre-ticked
    return Suggestion("family", old, new, "Probably a typo" if similar else "GBIF uses this family", similar)


class NameChecker(QObject):
    """Checks every distinct species name in parallel; emits finished(suggestions, problems)."""
    progress = pyqtSignal(int, int)
    finished = pyqtSignal(list, list)

    def __init__(self, user_agent, parent=None):
        super().__init__(parent)
        self.network = QNetworkAccessManager(self)
        self.user_agent = user_agent
        self.replies = []

    def start(self, species_to_families, kingdom):
        """species_to_families: {species name: set of family values used with it}."""
        self.pending = dict(species_to_families)
        self.total = len(self.pending)
        self.done = 0
        self.suggestions, self.problems = [], []
        self.family_votes = {}  # family in the table -> {GBIF family: [species]}
        if not self.pending:
            self.finished.emit([], [])
            return
        for name in self.pending:
            url = QUrl(f"{MATCH_URL}?{urlencode({'scientificName': name, 'kingdom': kingdom})}")
            request = QNetworkRequest(url)
            request.setHeader(QNetworkRequest.KnownHeaders.UserAgentHeader, self.user_agent)
            request.setTransferTimeout(20000)
            reply = self.network.get(request)
            reply.finished.connect(lambda reply=reply, name=name: self.on_reply(reply, name))
            self.replies.append(reply)

    def cancel(self):
        # Only replies still running: finished ones are already scheduled for deletion
        for reply in list(self.replies):
            reply.abort()

    def on_reply(self, reply, name):
        try:
            if reply.error() == QNetworkReply.NetworkError.OperationCanceledError:
                return
            if reply.error() != QNetworkReply.NetworkError.NoError:
                self.problems.append((name, f"Network error: {reply.errorString()}"))
            else:
                suggestion, family, problem = interpret_match(name, json.loads(bytes(reply.readAll())))
                if suggestion:
                    self.suggestions.append(suggestion)
                if problem:
                    self.problems.append((name, problem))
                if family:
                    for old_family in self.pending[name]:
                        self.family_votes.setdefault(old_family, {}).setdefault(family, []).append(name)
        except ValueError:
            self.problems.append((name, "Unexpected answer from GBIF"))
        finally:
            self.replies.remove(reply)
            reply.deleteLater()
            self.done += 1
            self.progress.emit(self.done, self.total)
            if self.done == self.total:
                self.finished.emit(self.suggestions + self.family_suggestions(), self.problems)

    def family_suggestions(self):
        suggestions = []
        for old, votes in self.family_votes.items():
            if len(votes) == 1:
                change = family_suggestion(old, next(iter(votes)))
                if change:
                    suggestions.append(change)
                continue
            # The same family name is used for species GBIF puts in different families:
            # renaming it everywhere would break the correct rows, so only report it
            for family, species in votes.items():
                if family != old:
                    for name in species:
                        self.problems.append((name, f"Family {old} in the table, GBIF says {family}"))
        return suggestions


def names_to_check(df, columns):
    species_column, family_column = columns.get("species"), columns.get("family")
    result = {}
    if not species_column:
        return result
    for _, row in df.iterrows():
        name = str(filter_nan(row[species_column])).strip()
        if not name:
            continue
        family = str(filter_nan(row[family_column])).strip() if family_column else ""
        result.setdefault(name, set())
        if family:
            result[name].add(family)
    return result


def save_corrected_copy(source, target, columns, suggestions):
    """Apply suggestions to a copy of the workbook, keeping its formatting. Returns the number of changed cells."""
    workbook = load_workbook(source)
    sheet = workbook.worksheets[0]
    header = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
    changed = 0
    for field in ("species", "family"):
        mapping = {s.old: s.new for s in suggestions if s.field == field}
        column = header.get(columns.get(field) or "")
        if not mapping or not column:
            continue
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            if isinstance(cell.value, str) and cell.value.strip() in mapping:
                cell.value = mapping[cell.value.strip()]
                changed += 1
    workbook.save(target)
    return changed


def to_darwin_core(df, columns, herbarium, herbarium_type):
    """One Darwin Core occurrence per row, as a DataFrame ready to be saved as CSV."""
    records = []
    for _, row in df.iterrows():
        get = lambda field: filter_nan(row.get(columns[field], '')) if columns.get(field) else ''
        num = format_num(get("num")).strip()
        lat, lon = parse_coordinate(get("N")), parse_coordinate(get("E"))
        records.append({
            "occurrenceID": f"{herbarium}:{num}" if num else "",
            "basisOfRecord": "PreservedSpecimen",
            "institutionCode": herbarium,
            "catalogNumber": num,
            "scientificName": get("species"),
            "kingdom": kingdom_for(herbarium_type),
            "family": get("family"),
            "vernacularName": get("speciesrus"),
            "eventDate": remove_time_from_date(get("date")),
            "recordedBy": get("leg."),
            "identifiedBy": get("det."),
            "stateProvince": get("region"),
            "locality": get("point"),
            "habitat": get("habitats"),
            "decimalLatitude": lat if lat is not None else "",
            "decimalLongitude": lon if lon is not None else "",
            "geodeticDatum": "WGS84" if lat is not None and lon is not None else "",
            "associatedReferences": inat_url(get("inat")),
        })
    return pd.DataFrame(records)
